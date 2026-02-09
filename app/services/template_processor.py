import re
import json
from typing import Dict, List, Optional, Any
import tempfile
import os

import aiofiles
from openai import AsyncOpenAI

from app.core.config import get_settings

settings = get_settings()


class TemplateProcessor:
    """Service for processing sample documents and converting them to templates using LLM."""
    
    def __init__(self, openai_client: Optional[AsyncOpenAI] = None):
        if openai_client:
            self.openai_client = openai_client
        else:
            # Only create OpenAI client if API key is available
            if settings.openai_api_key:
                self.openai_client = AsyncOpenAI(api_key=settings.openai_api_key)
            else:
                self.openai_client = None
        self.supported_extensions = {'.docx', '.doc', '.pdf', '.txt', '.md'}
    
    async def convert_sample_to_template(
        self,
        file_content: bytes,
        filename: str,
        template_name: str,
        version: str = "1.0",
        kind: str = "document",
        template_description: Optional[str] = None,
    ) -> Dict[str, Any]:
        """
        Convert a sample document into a reusable template using LLM.
        
        The sample document is treated as an example that shows the desired structure
        and format. The LLM identifies what content should be replaced with variables
        to make it reusable for generating documents from suggestion clusters.
        
        Args:
            file_content: Binary content of the uploaded sample document
            filename: Original filename 
            template_name: Name for the template
            version: Template version
            kind: Template kind/type
            template_description: Optional description for the template
            
        Returns:
            Dictionary containing template data ready for creation
        """
        file_ext = self._get_file_extension(filename)
        
        if file_ext not in self.supported_extensions:
            raise ValueError(f"Unsupported file format: {file_ext}")
        
        # Extract text content from sample document
        sample_content = await self._extract_text_content(file_content, file_ext)
        
        if not sample_content.strip():
            raise ValueError("No readable content found in the sample document")
        
        # Use LLM to convert sample to template
        template_data = await self._convert_sample_with_llm(
            sample_content=sample_content,
            filename=filename,
            template_name=template_name,
            kind=kind,
            template_description=template_description,
        )
        
        return {
            "name": template_name,
            "description": template_description or template_data.get("description", f"Template created from sample: {filename}"),
            "version": version,
            "kind": kind,
            "content_markdown": template_data["content_markdown"],
            "outline": template_data.get("outline", []),
            "placeholders": template_data.get("placeholders", []),
        }
    
    def _get_file_extension(self, filename: str) -> str:
        """Extract file extension from filename."""
        return os.path.splitext(filename.lower())[1]
    
    async def _extract_text_content(self, file_content: bytes, file_ext: str) -> str:
        """Extract text from various file formats."""
        if file_ext == '.pdf':
            return await self._extract_from_pdf(file_content)
        elif file_ext in ['.docx', '.doc']:
            return await self._extract_from_word(file_content)
        elif file_ext in ['.txt', '.md']:
            return file_content.decode('utf-8', errors='ignore')
        else:
            raise ValueError(f"Unsupported file format: {file_ext}")
    
    async def _extract_from_pdf(self, file_content: bytes) -> str:
        """Extract text from PDF file."""
        try:
            import PyPDF2
            
            with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as temp_file:
                temp_file.write(file_content)
                temp_file_path = temp_file.name
            
            try:
                with open(temp_file_path, 'rb') as file:
                    reader = PyPDF2.PdfReader(file)
                    text_content = ""
                    
                    for page in reader.pages:
                        text_content += page.extract_text() + "\n"
                
                return text_content.strip()
            finally:
                os.unlink(temp_file_path)
                
        except ImportError:
            raise ValueError("PyPDF2 not installed. Cannot process PDF files.")
        except Exception as e:
            raise ValueError(f"Failed to process PDF: {str(e)}")
    
    async def _extract_from_word(self, file_content: bytes) -> str:
        """Extract text from Word document."""
        try:
            import docx
            
            with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as temp_file:
                temp_file.write(file_content)
                temp_file_path = temp_file.name
            
            try:
                doc = docx.Document(temp_file_path)
                text_content = ""
                
                for paragraph in doc.paragraphs:
                    text_content += paragraph.text + "\n"
                
                return text_content.strip()
            finally:
                os.unlink(temp_file_path)
                
        except ImportError:
            raise ValueError("python-docx not installed. Cannot process Word documents.")
        except Exception as e:
            raise ValueError(f"Failed to process Word document: {str(e)}")
    
    async def _convert_sample_with_llm(
        self,
        sample_content: str,
        filename: str,
        template_name: str,
        kind: str,
        template_description: Optional[str] = None,
    ) -> Dict[str, Any]:
        """Use LLM to convert sample document content to a reusable template."""
        
        system_prompt = """You are an expert at converting sample documents into reusable templates for a digital suggestion box system. 

Your task is to analyze a sample document and convert it into a template that can be used to generate similar documents from suggestion clusters and data.

CONTEXT: This template will be used to generate documents from grouped suggestions in a digital suggestion box. The suggestions contain information like:
- Titles and descriptions of suggestions
- Categories (UX, Product, Service, Operational, Other)  
- Author types (Staff, Customer)
- Tags and topics
- Contact information
- Timestamps

INSTRUCTIONS:
1. Convert the sample document to clean, well-structured markdown
2. Identify content that should become dynamic variables based on suggestion data
3. Replace specific values with meaningful placeholders in {{variable_name}} format
4. Create an outline showing the document structure
5. Generate a list of all placeholders used

PLACEHOLDER GUIDELINES:
- Use descriptive names like: {{cluster_title}}, {{suggestion_count}}, {{primary_topic}}, {{key_themes}}
- Consider what data would come from suggestion clusters: themes, topics, categories, summaries
- Replace specific dates with {{report_date}} or {{analysis_date}}
- Replace names with {{contact_name}}, {{department}}, {{team_name}}
- Replace metrics with {{total_suggestions}}, {{response_rate}}, etc.
- Replace specific content with {{executive_summary}}, {{key_findings}}, {{recommendations}}

Return ONLY valid JSON with this structure:
{
  "description": "Brief description of what this template generates",
  "content_markdown": "Markdown content with {{placeholders}}",
  "outline": [{"level": 1, "text": "Section Title", "line": 5}],
  "placeholders": ["variable1", "variable2"]
}"""

        user_prompt = f"""Convert this sample document into a reusable template:

**Sample Document:** {filename}
**Template Name:** {template_name}  
**Template Type:** {kind}
**Description:** {template_description or 'Not provided'}

**Sample Content:**
{sample_content}

Analyze this sample and convert it into a template that can generate similar documents using data from suggestion clusters. Focus on identifying what should be variable vs what should remain static."""

        try:
            response = await self.openai_client.chat.completions.create(
                model="gpt-4o-mini",
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_prompt}
                ],
                temperature=0.2,  # Lower temperature for more consistent output
                max_tokens=4000,
            )
            
            response_content = response.choices[0].message.content.strip()
            
            # Extract and parse JSON response
            template_data = self._parse_llm_response(response_content)
            
            # Validate and enhance the response
            template_data = self._validate_template_data(template_data, filename)
            
            return template_data
            
        except Exception as e:
            raise ValueError(f"Failed to convert sample to template with LLM: {str(e)}")
    
    def _parse_llm_response(self, response_content: str) -> Dict[str, Any]:
        """Parse JSON response from LLM, handling various formats."""
        # Try to find JSON in the response
        json_patterns = [
            r'\{.*\}',  # Basic JSON pattern
            r'```json\s*(\{.*\})\s*```',  # JSON in code blocks
            r'```\s*(\{.*\})\s*```',  # JSON in unmarked code blocks
        ]
        
        for pattern in json_patterns:
            match = re.search(pattern, response_content, re.DOTALL)
            if match:
                json_content = match.group(1) if match.groups() else match.group(0)
                try:
                    return json.loads(json_content)
                except json.JSONDecodeError:
                    continue
        
        # If no valid JSON found, create fallback structure
        return {
            "description": "Template created from sample document",
            "content_markdown": self._clean_response_content(response_content),
            "outline": [],
            "placeholders": []
        }
    
    def _clean_response_content(self, content: str) -> str:
        """Clean up content that's not in JSON format."""
        # Remove code block markers if present
        content = re.sub(r'^```[a-zA-Z]*\n?', '', content)
        content = re.sub(r'\n?```$', '', content)
        return content.strip()
    
    def _validate_template_data(self, template_data: Dict[str, Any], filename: str) -> Dict[str, Any]:
        """Validate and enhance template data from LLM."""
        # Ensure required fields exist
        if "content_markdown" not in template_data or not template_data["content_markdown"]:
            raise ValueError("Template must contain markdown content")
        
        # Set default description if missing
        if not template_data.get("description"):
            template_data["description"] = f"Template created from sample: {filename}"
        
        # Extract outline if missing or empty
        if not template_data.get("outline"):
            template_data["outline"] = self._extract_outline_from_markdown(
                template_data["content_markdown"]
            )
        
        # Extract placeholders if missing or empty
        if not template_data.get("placeholders"):
            template_data["placeholders"] = self._extract_placeholders_from_markdown(
                template_data["content_markdown"]
            )
        
        return template_data
    
    def _extract_outline_from_markdown(self, markdown_content: str) -> List[Dict[str, Any]]:
        """Extract document outline from markdown headers."""
        lines = markdown_content.split('\n')
        outline = []
        
        for line_num, line in enumerate(lines, 1):
            line = line.strip()
            if line.startswith('#'):
                level = len(line) - len(line.lstrip('#'))
                heading_text = line.lstrip('#').strip()
                
                if heading_text:  # Only add non-empty headings
                    outline.append({
                        "level": level,
                        "text": heading_text,
                        "line": line_num,
                    })
        
        return outline
    
    def _extract_placeholders_from_markdown(self, markdown_content: str) -> List[str]:
        """Extract all placeholders from markdown content."""
        # Find all {{placeholder}} patterns
        placeholder_pattern = r'\{\{([^}]+)\}\}'
        matches = re.findall(placeholder_pattern, markdown_content)
        
        placeholders = []
        for match in matches:
            placeholder = match.strip()
            if placeholder and placeholder not in placeholders:
                placeholders.append(placeholder)
        
        return sorted(placeholders)


# Global instance
template_processor = TemplateProcessor()