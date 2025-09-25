import io
import tempfile
import os
from typing import Optional, BinaryIO
from uuid import UUID

import markdown
from weasyprint import HTML, CSS
from weasyprint.text.fonts import FontConfiguration
import aiofiles

from app.domain.models import Document
from app.domain.enums import DocFormat


class DocumentExporter:
    """Service for exporting documents to various formats."""
    
    def __init__(self):
        self.font_config = FontConfiguration()
        
        # Default CSS for PDF export
        self.default_pdf_css = """
        @page {
            size: A4;
            margin: 2cm;
        }
        
        body {
            font-family: 'Helvetica', 'Arial', sans-serif;
            font-size: 11pt;
            line-height: 1.4;
            color: #333;
            margin: 0;
            padding: 0;
        }
        
        h1, h2, h3, h4, h5, h6 {
            color: #2c3e50;
            margin-top: 1.5em;
            margin-bottom: 0.5em;
            font-weight: 600;
        }
        
        h1 {
            font-size: 1.8em;
            border-bottom: 2px solid #3498db;
            padding-bottom: 0.3em;
        }
        
        h2 {
            font-size: 1.4em;
            border-bottom: 1px solid #bdc3c7;
            padding-bottom: 0.2em;
        }
        
        h3 {
            font-size: 1.2em;
        }
        
        p {
            margin-bottom: 1em;
            text-align: justify;
        }
        
        ul, ol {
            margin-bottom: 1em;
            padding-left: 2em;
        }
        
        li {
            margin-bottom: 0.3em;
        }
        
        blockquote {
            border-left: 4px solid #3498db;
            margin: 1em 0;
            padding: 0.5em 0 0.5em 1em;
            background-color: #f8f9fa;
            font-style: italic;
        }
        
        code {
            background-color: #f1f2f6;
            padding: 0.2em 0.4em;
            border-radius: 3px;
            font-family: 'Courier New', monospace;
            font-size: 0.9em;
        }
        
        pre {
            background-color: #f1f2f6;
            padding: 1em;
            border-radius: 5px;
            overflow-x: auto;
            font-family: 'Courier New', monospace;
            font-size: 0.9em;
        }
        
        table {
            border-collapse: collapse;
            width: 100%;
            margin-bottom: 1em;
        }
        
        th, td {
            border: 1px solid #bdc3c7;
            padding: 0.5em;
            text-align: left;
        }
        
        th {
            background-color: #ecf0f1;
            font-weight: 600;
        }
        
        .header {
            text-align: center;
            margin-bottom: 2em;
        }
        
        .footer {
            margin-top: 2em;
            padding-top: 1em;
            border-top: 1px solid #bdc3c7;
            font-size: 0.9em;
            color: #7f8c8d;
            text-align: center;
        }
        """

    async def export_document(
        self, 
        document: Document, 
        export_format: DocFormat,
        custom_css: Optional[str] = None
    ) -> bytes:
        """Export document to specified format."""
        
        if not document.draft_content or 'markdown' not in document.draft_content:
            raise ValueError("Document has no markdown content to export")
        
        markdown_content = document.draft_content['markdown']
        
        if export_format == DocFormat.PDF:
            return await self._export_to_pdf(markdown_content, document, custom_css)
        elif export_format == DocFormat.DOCX:
            return await self._export_to_docx(markdown_content, document)
        else:
            raise ValueError(f"Unsupported export format: {export_format}")

    async def export_to_text(self, document: Document) -> str:
        """Export document as plain text."""
        if not document.draft_content or 'markdown' not in document.draft_content:
            raise ValueError("Document has no markdown content to export")
        
        markdown_content = document.draft_content['markdown']
        
        # Convert markdown to plain text by removing markdown syntax
        import re
        
        # Remove headers
        text = re.sub(r'^#{1,6}\s+', '', markdown_content, flags=re.MULTILINE)
        
        # Remove bold/italic
        text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
        text = re.sub(r'\*(.*?)\*', r'\1', text)
        
        # Remove links but keep text
        text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', text)
        
        # Remove code blocks
        text = re.sub(r'```[^`]*```', '', text, flags=re.DOTALL)
        text = re.sub(r'`([^`]+)`', r'\1', text)
        
        # Clean up multiple newlines
        text = re.sub(r'\n\s*\n', '\n\n', text)
        
        return text.strip()

    async def export_to_markdown(self, document: Document) -> str:
        """Export document as markdown."""
        if not document.draft_content or 'markdown' not in document.draft_content:
            raise ValueError("Document has no markdown content to export")
        
        return document.draft_content['markdown']

    async def _export_to_pdf(
        self, 
        markdown_content: str, 
        document: Document,
        custom_css: Optional[str] = None
    ) -> bytes:
        """Export markdown content to PDF."""
        
        # Convert markdown to HTML
        html_content = markdown.markdown(
            markdown_content,
            extensions=[
                'extra',      # Tables, fenced code blocks, etc.
                'codehilite', # Syntax highlighting
                'toc',        # Table of contents
                'sane_lists', # Better list handling
            ]
        )
        
        # Wrap in proper HTML structure
        full_html = f"""
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="utf-8">
            <title>{document.title or 'Document'}</title>
        </head>
        <body>
            <div class="header">
                <h1>{document.title or 'Document'}</h1>
                <p><em>Generated on {document.updated_at.strftime('%Y-%m-%d %H:%M UTC')}</em></p>
            </div>
            
            {html_content}
            
            <div class="footer">
                <p>Document ID: {document.id}</p>
                <p>Generated by Digital Suggestion Box</p>
            </div>
        </body>
        </html>
        """
        
        # Use custom CSS or default
        css_content = custom_css or self.default_pdf_css
        
        try:
            # Generate PDF
            html_doc = HTML(string=full_html)
            css_doc = CSS(string=css_content, font_config=self.font_config)
            
            # Create PDF in memory
            pdf_bytes = html_doc.write_pdf(stylesheets=[css_doc], font_config=self.font_config)
            return pdf_bytes
            
        except Exception as e:
            raise ValueError(f"Failed to generate PDF: {str(e)}")

    async def _export_to_docx(self, markdown_content: str, document: Document) -> bytes:
        """Export markdown content to DOCX."""
        
        try:
            import pypandoc
            
            # Create temporary markdown file
            with tempfile.NamedTemporaryFile(mode='w', suffix='.md', delete=False, encoding='utf-8') as md_file:
                md_file.write(markdown_content)
                md_file_path = md_file.name
            
            try:
                # Convert to DOCX using pandoc
                docx_content = pypandoc.convert_file(
                    md_file_path,
                    'docx',
                    format='markdown',
                    extra_args=[
                        '--standalone',
                        f'--metadata=title:{document.title or "Document"}',
                        '--reference-doc=word-template.docx' if os.path.exists('word-template.docx') else '',
                    ]
                )
                
                return docx_content
                
            finally:
                # Clean up temporary file
                os.unlink(md_file_path)
                
        except ImportError:
            # Fallback: use python-docx for basic conversion
            return await self._export_to_docx_basic(markdown_content, document)
        except Exception as e:
            raise ValueError(f"Failed to generate DOCX: {str(e)}")

    async def _export_to_docx_basic(self, markdown_content: str, document: Document) -> bytes:
        """Basic DOCX export using python-docx (fallback method)."""
        
        try:
            from docx import Document as DocxDocument
            from docx.shared import Inches
            import re
            
            # Create new document
            doc = DocxDocument()
            
            # Add title
            title = doc.add_heading(document.title or 'Document', 0)
            
            # Add metadata
            doc.add_paragraph(f'Generated on {document.updated_at.strftime("%Y-%m-%d %H:%M UTC")}').italic = True
            doc.add_paragraph('')  # Empty line
            
            # Parse markdown content (basic parsing)
            lines = markdown_content.split('\n')
            
            for line in lines:
                line = line.strip()
                
                if not line:
                    doc.add_paragraph('')  # Empty paragraph
                elif line.startswith('# '):
                    doc.add_heading(line[2:], level=1)
                elif line.startswith('## '):
                    doc.add_heading(line[3:], level=2)
                elif line.startswith('### '):
                    doc.add_heading(line[4:], level=3)
                elif line.startswith('#### '):
                    doc.add_heading(line[5:], level=4)
                elif line.startswith('- ') or line.startswith('* '):
                    # Bullet point
                    para = doc.add_paragraph(line[2:], style='List Bullet')
                elif re.match(r'^\d+\. ', line):
                    # Numbered list
                    para = doc.add_paragraph(line[3:], style='List Number')
                else:
                    # Regular paragraph
                    para = doc.add_paragraph(line)
            
            # Save to bytes
            doc_bytes = io.BytesIO()
            doc.save(doc_bytes)
            doc_bytes.seek(0)
            
            return doc_bytes.getvalue()
            
        except ImportError:
            raise ValueError("python-docx not available for DOCX export")
        except Exception as e:
            raise ValueError(f"Failed to generate basic DOCX: {str(e)}")

    def get_content_type(self, format: DocFormat) -> str:
        """Get MIME content type for format."""
        if format == DocFormat.PDF:
            return 'application/pdf'
        elif format == DocFormat.DOCX:
            return 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        else:
            raise ValueError(f"Unknown format: {format}")

    def get_file_extension(self, format: DocFormat) -> str:
        """Get file extension for format."""
        if format == DocFormat.PDF:
            return '.pdf'
        elif format == DocFormat.DOCX:
            return '.docx'
        else:
            raise ValueError(f"Unknown format: {format}")


# Global instance
document_exporter = DocumentExporter()